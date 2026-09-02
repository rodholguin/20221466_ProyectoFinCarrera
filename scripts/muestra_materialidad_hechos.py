"""Muestra para medir la TASA DE MATERIALIDAD de los Hechos de Importancia BVL.

LA PREGUNTA QUE RESPONDE. La evaluación integral de R5
(docs/_TEMP_para_decidir_20260829.txt §2) propone sumar el feed oficial de
Hechos de Importancia como SEGUNDA FUENTE DE EVENTOS, para los 7 activos, junto
a la prensa —no en vez de ella—. El argumento es que nivela la cobertura entre
activos (razón max/min 16.8x -> 2.9x) porque todo emisor listado está obligado a
divulgar, mientras la prensa cubre a quien quiere. Pero ese cálculo descansa en
una estimación GENEROSA: que los 1,549 "candidatos" que quedan tras filtrar
rutina son eventos de verdad. NO ESTÁ MEDIDO. Esto lo mide.

POR QUÉ NO LO ETIQUETA EL ASISTENTE. Es la lección del BLOQUEANTE #1 de R5: la
referencia de julio la produjo el asistente y sesgó la comparación entera de
modelos a favor de los que razonaban como él. La referencia la hace el autor, a
ciegas, o no vale.

DISEÑO DE LA MUESTRA
  - Estrato = ticker x {candidato, rutina}. Los 7 activos entran sí o sí: el
    interés está justamente en los que la prensa no cubre (Ferreycorp, Pacasmayo).
  - El estrato RUTINA no es relleno: es la VERIFICACIÓN DEL FILTRO, igual que se
    hizo con el prefiltro regex de R5 ("descarta 19.8% con CERO relevantes
    perdidos"). Si aparecen materiales dentro de la rutina, el filtro está mal y
    hay que aflojarlo.
  - Pesos N_estrato/n_estrato -> estimador de Hájek, como en
    scripts/evalua_validacion_humana_r5.py. El promedio simple NO es el de la
    población: la muestra sobre-representa a los activos chicos a propósito.

QUÉ VE EL ANOTADOR. `observation` + el CÓDIGO OFICIAL del hecho. Los dos, porque
los dos van a estar disponibles en producción y point-in-time. Ocultar el código
mediría una tarea más difícil que la real. NO ve la clasificación
rutina/candidato del asistente, que es lo que se está poniendo a prueba.

Uso:
  python scripts/muestra_materialidad_hechos.py
  python scripts/muestra_materialidad_hechos.py --n-candidatos 15 --n-rutina 5
"""
from __future__ import annotations

import argparse
import ast
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from src.sentiment.taxonomia import TAXONOMIA  # noqa: E402

CACHE = ROOT / "data" / "raw" / "smv_cache"
SALIDA = ROOT / "data" / "interim" / "validacion_humana"
SEMILLA = 20221466

AZUL, AMARILLO, GRIS, FUENTE = "1F3864", "FFF2CC", "EEEEEE", "Calibri"

# rpj -> ticker. Los rpj salen de config.yaml (campo smv_rpj).
RPJ = {"B80005": "CREDITC1", "A20032": "MINSURI1", "B30006": "ALICORC1",
       "OE5087": "INRETC1", "CD0005": "CPACASC1", "B60001": "FERREYC1",
       "B40008": "LUSURC1"}

# ── FILTRO DE RUTINA ─────────────────────────────────────────────────────────
# Códigos oficiales que son divulgación DE CALENDARIO: existen porque la norma
# obliga a publicarlas cada cierto tiempo, no porque haya pasado algo.
RUT_CODE = {
    "E05",  # Aprobación de información financiera anual auditada / memoria
    "L01", "E01",  # Convocatoria a juntas
    "L40",  # Posición mensual en instrumentos financieros derivados
    "B03",  # Anuncio y resultados de colocación; cronograma y pagos
    "Z96",  # Expediente de listado o deslistado
    "L11",  # Designación de sociedad de auditoría
    "E02",  # Resultado de acuerdos de junta
    "L02", "L03", "E04", "E03",
}
# Rutina detectable por TEXTO, para lo que cae en códigos genéricos (L37 "Otros"
# es el 34% del feed y ahí adentro hay EEFF mensuales y reportes periódicos de
# ADRs). Cada patrón describe una publicación PERIÓDICA, no un hecho.
RUT_TXT = re.compile(
    r"(ee\.?\s?ff|estados?\s+financieros?|informaci[oó]n\s+financiera|"
    r"al\s+\d{1,2}[-/ ](ene|feb|mar|abr|may|jun|jul|ago|set|sep|oct|nov|dic)|"
    r"posici[oó]n\s+mensual|memoria\s+anual|"
    r"tenedores\s+de\s+adrs?|transacciones\s+efectuadas\s+con\s+adrs?|"
    r"c[oó]digo\s+de\s+buen\s+gobierno)", re.I)


def cargar() -> pd.DataFrame:
    """Lee el caché de Hechos de Importancia de los 7 activos y lo aplana."""
    filas = []
    for rpj, tic in RPJ.items():
        f = CACHE / f"bvl_hechos_{rpj}_2012_2025.json"
        if not f.exists():
            raise FileNotFoundError(
                f"falta {f}. Descargar con "
                f"src.fundamentals.fundamentals_client._bvl_fetch_hechos"
                f"('{rpj}', '2012-01-01', '2025-12-31', cache_dir)")
        for x in json.loads(f.read_text(encoding="utf-8")):
            cods = x.get("codes")
            if isinstance(cods, str):
                try:
                    cods = ast.literal_eval(cods)
                except (ValueError, SyntaxError):
                    cods = []
            cods = cods or [{}]
            filas.append({
                "ticker": tic,
                "empresa": x.get("businessName"),
                # registerDate = momento de DIVULGACIÓN. Es la fecha
                # point-in-time correcta; sessionDate puede ser anterior.
                "fecha": str(x.get("registerDate"))[:10],
                "codigo": cods[0].get("codeHHII") or "",
                # El feed trae el mismo código con distinta capitalización.
                "desc_codigo": str(cods[0].get("descCodeHHII") or "").strip(),
                "hecho": str(x.get("observation") or "").strip(),
            })
    d = pd.DataFrame(filas)
    d["anio"] = d.fecha.str[:4].astype(int)
    d = d[(d.anio >= 2012) & (d.anio <= 2025)].copy()
    d["desc_codigo"] = d.desc_codigo.str.capitalize()

    d["_vacia"] = d.hecho.str.len() < 8
    d["_rutina"] = (d.codigo.isin(RUT_CODE)
                    | d.hecho.str.contains(RUT_TXT, regex=True))
    d["capa"] = "candidato"
    d.loc[d._rutina, "capa"] = "rutina"
    d.loc[d._vacia, "capa"] = "sin_texto"
    return d.reset_index(drop=True)


def construir(n_cand: int, n_rut: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    d = cargar()
    print(f"población: {len(d)} hechos de importancia (7 activos, 2012-2025)")
    print(d.capa.value_counts().to_string())
    print("\n  candidatos por activo:")
    print(d[d.capa == "candidato"].groupby("ticker").size().to_string())

    # `sin_texto` se excluye del muestreo: no es anotable por texto. Se cuenta y
    # se declara como límite (2,007 filas), no se finge que se midió.
    pool = d[d.capa != "sin_texto"].copy()
    pool["estrato"] = pool.ticker + "|" + pool.capa

    partes = []
    for est, g in pool.groupby("estrato"):
        k = n_cand if est.endswith("candidato") else n_rut
        k = min(k, len(g))
        partes.append(g.sample(k, random_state=SEMILLA))
    m = pd.concat(partes, ignore_index=True)

    tam = pool.groupby("estrato").size().rename("n_pool")
    mue = m.groupby("estrato").size().rename("n_muestra")
    pesos = pd.concat([tam, mue], axis=1)
    pesos["peso"] = pesos.n_pool / pesos.n_muestra
    m["peso"] = m.estrato.map(pesos.peso)

    print(f"\nmuestra: {len(m)} filas  ({m.capa.value_counts().to_dict()})")
    print(pesos.to_string())

    # Orden aleatorio: así cualquier PREFIJO de la anotación es una submuestra
    # válida y el autor puede parar cuando quiera sin sesgar nada.
    m = m.sample(frac=1.0, random_state=SEMILLA + 1).reset_index(drop=True)
    m.insert(0, "n", range(1, len(m) + 1))

    clave = m[["n", "ticker", "capa", "estrato", "peso", "codigo", "fecha"]]
    return m, clave


# ─────────────────────────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────────────────────────
def _taxonomia() -> list[tuple[str, float, str, str]]:
    """Misma taxonomía que R5, con el prefijo A/B/C/Z por magnitud.

    Es DELIBERADO usar la misma: el objetivo es que las dos fuentes de eventos
    entren por la misma escala y el mismo encoding de D15. Si hiciera falta una
    taxonomía distinta para los hechos, eso ya sería un resultado.
    """
    pref = {1.0: "A", 0.6: "B", 0.3: "C", 0.0: "Z"}
    return [(f"{pref[m]}_{nom}", m, nom, d) for nom, (m, d) in TAXONOMIA.items()]


def escribir_excel(hoja: pd.DataFrame, destino: Path) -> None:
    taxonomia = _taxonomia()
    wb = Workbook()

    ins = wb.active
    ins.title = "Instrucciones"
    ins.sheet_view.showGridLines = False
    ins.column_dimensions["A"].width = 4
    ins.column_dimensions["B"].width = 108

    filas = [
        ("t", "Materialidad de los Hechos de Importancia (BVL) — muestra a ciegas"),
        ("", ""),
        ("p", "Esto decide si el feed oficial de Hechos de Importancia entra como SEGUNDA "
              "fuente de eventos del canal de sentimiento, JUNTO a la prensa y para los 7 "
              "activos. La prensa no cubre a media cartera (Ferreycorp: 0.9 eventos con "
              "signo por ano); los hechos si, porque todo emisor listado esta obligado a "
              "divulgar. Lo que falta medir es cuantos de esos hechos son EVENTOS DE VERDAD "
              "y cuantos son tramite."),
        ("", ""),
        ("h", "LA PREGUNTA, EN UNA LINEA"),
        ("p", "Para cada fila: si usted fuera accionista de esa empresa y leyera ese hecho el "
              "dia que se publico, ¿le cambia algo? Si si, elija la categoria A/B/C que "
              "corresponda. Si es tramite, papeleo o rutina, elija una Z."),
        ("", ""),
        ("h", "QUE se llena (dos columnas amarillas)"),
        ("p", "1) categoria — SIEMPRE. Es la MISMA taxonomia de la anotacion de prensa, a "
              "proposito: las dos fuentes tienen que entrar por la misma escala."),
        ("p", "2) polaridad — SOLO si la categoria NO empieza con Z."),
        ("", ""),
        ("h", "LO QUE VA A CHIRRIAR, Y ESTA BIEN QUE CHIRRIE"),
        ("p", "Muchos hechos son emisiones de bonos, actas de directorio o prospectos. La "
              "taxonomia de prensa no fue disenada para eso. Si ninguna categoria calza, "
              "elija la menos mala, marque 'duda' y digalo en 'nota': que la taxonomia no "
              "sirva para esta fuente es un resultado, no un error suyo."),
        ("", ""),
        ("h", "OJO CON EL DOBLE CONTEO"),
        ("p", "Los dividendos ya estan en el panel desde R3 y los EEFF desde R4. Si el hecho "
              "es exactamente eso, marquelo igual segun la taxonomia y anote 'ya en panel' "
              "en la nota: se excluiran despues para no contarlos dos veces."),
        ("", ""),
        ("h", "COLUMNA `codigo`"),
        ("p", "Es el codigo oficial que el emisor le puso al hecho. Uselo: va a estar "
              "disponible en produccion igual que ahora. 'L37 Otros hechos de importancia' "
              "es el 34% del feed y no dice nada — ahi hay que leer el texto."),
        ("", ""),
        ("h", "PUEDE PARAR CUANDO QUIERA"),
        ("p", "Las filas estan en orden ALEATORIO, asi que cualquier prefijo es una muestra "
              "valida. No ordenar ni filtrar de forma permanente: el orden es la clave."),
        ("", ""),
        ("h", "AL TERMINAR"),
        ("p", "Guarde con el mismo nombre y avise."),
        ("", ""),
        ("h", "PROGRESO"),
    ]
    r = 1
    for tipo, txt in filas:
        if tipo == "t":
            ins[f"B{r}"] = txt
            ins[f"B{r}"].font = Font(name=FUENTE, size=15, bold=True, color=AZUL)
        elif tipo == "h":
            ins[f"B{r}"] = txt
            ins[f"B{r}"].font = Font(name=FUENTE, size=11, bold=True, color=AZUL)
        elif tipo == "p":
            ins[f"B{r}"] = txt
            ins[f"B{r}"].font = Font(name=FUENTE, size=10)
            ins[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            ins.row_dimensions[r].height = 15 * (1 + len(txt) // 105)
        r += 1

    n = len(hoja)
    r0 = r
    prog = [
        ("Filas en total", n),
        ("Anotadas (categoria llena)", f"=COUNTA(Anotacion!$F$2:$F${n + 1})"),
        ("Avance", f"=C{r0 + 1}/C{r0}"),
        ("Materiales hasta ahora (A/B/C)",
         f'=COUNTA(Anotacion!$F$2:$F${n + 1})'
         f'-COUNTIF(Anotacion!$F$2:$F${n + 1},"Z_*")'),
        ("Marcadas con duda", f'=COUNTIF(Anotacion!$H$2:$H${n + 1},"x")'),
    ]
    for etiqueta, valor in prog:
        ins[f"B{r}"] = etiqueta
        ins[f"C{r}"] = valor
        ins[f"B{r}"].font = Font(name=FUENTE, size=10)
        ins[f"C{r}"].font = Font(name=FUENTE, size=10, bold=True)
        r += 1
    ins.column_dimensions["C"].width = 14
    ins[f"C{r0 + 2}"].number_format = "0.0%"

    # ---------------- Taxonomia ----------------
    tx = wb.create_sheet("Taxonomia")
    tx.sheet_view.showGridLines = False
    for j, c in enumerate(["codigo", "magnitud", "nombre_taxonomia", "definicion"], 1):
        cel = tx.cell(row=1, column=j, value=c)
        cel.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", start_color=AZUL)
    for i, (cod, mag, nom, defi) in enumerate(taxonomia, start=2):
        tx.cell(row=i, column=1, value=cod).font = Font(name=FUENTE, size=10, bold=True)
        tx.cell(row=i, column=2, value=mag).number_format = "0.0"
        tx.cell(row=i, column=3, value=nom).font = Font(name=FUENTE, size=10)
        c = tx.cell(row=i, column=4, value=defi)
        c.font = Font(name=FUENTE, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if mag == 0.0:
            for j in range(1, 5):
                tx.cell(row=i, column=j).fill = PatternFill("solid", start_color=GRIS)
    for col, w in zip("ABCD", (26, 10, 24, 78)):
        tx.column_dimensions[col].width = w
    tx.freeze_panes = "A2"
    tx["F1"] = "Las filas grises son de magnitud CERO: son la abstencion (no material)."
    tx["F1"].font = Font(name=FUENTE, size=10, italic=True)

    # ---------------- Anotacion ----------------
    ws = wb.create_sheet("Anotacion")
    cab = ["n", "fecha", "codigo", "empresa", "hecho",
           "categoria", "polaridad", "duda", "nota", "desc_codigo"]
    anchos = [6, 11, 8, 30, 92, 26, 12, 7, 34, 46]
    for j, (c, w) in enumerate(zip(cab, anchos), 1):
        cel = ws.cell(row=1, column=j, value=c)
        cel.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", start_color=AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 22

    borde = Side(style="thin", color="D9D9D9")
    for i, fila in enumerate(hoja.itertuples(index=False), start=2):
        vals = [fila.n, fila.fecha, fila.codigo, fila.empresa, fila.hecho,
                None, None, None, None, fila.desc_codigo]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10)
            c.border = Border(bottom=borde)
            c.alignment = Alignment(vertical="center", wrap_text=(j == 5),
                                    horizontal="center" if j in (1, 2, 3, 8) else "left")
            if j in (6, 7, 8, 9):
                c.fill = PatternFill("solid", start_color=AMARILLO)
            if j == 10:
                c.font = Font(name=FUENTE, size=8, color="808080")
        ws.row_dimensions[i].height = 28

    dv = DataValidation(type="list",
                        formula1=f"=Taxonomia!$A$2:$A${len(taxonomia) + 1}",
                        allow_blank=True, showDropDown=False)
    dv.error = "Elija una categoria del desplegable (hoja Taxonomia)."
    dv.errorTitle = "Categoria no valida"
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(hoja) + 1}")

    dvp = DataValidation(type="list", formula1='"positivo,negativo,neutral"',
                         allow_blank=True, showDropDown=False)
    dvp.error = "Solo positivo, negativo o neutral. Vacia si la categoria es Z."
    dvp.errorTitle = "Polaridad no valida"
    ws.add_data_validation(dvp)
    dvp.add(f"G2:G{len(hoja) + 1}")

    dvd = DataValidation(type="list", formula1='"x"', allow_blank=True,
                         showDropDown=False)
    ws.add_data_validation(dvd)
    dvd.add(f"H2:H{len(hoja) + 1}")

    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:J{len(hoja) + 1}"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n-candidatos", type=int, default=15,
                    help="filas por activo del estrato `candidato`")
    ap.add_argument("--n-rutina", type=int, default=5,
                    help="filas por activo del estrato `rutina` (verifica el filtro)")
    a = ap.parse_args()

    m, clave = construir(a.n_candidatos, a.n_rutina)
    xlsx = SALIDA / "muestra_hechos_materialidad.xlsx"
    csv = SALIDA / "clave_hechos_materialidad.csv"
    escribir_excel(m, xlsx)
    clave.to_csv(csv, index=False, encoding="utf-8")
    print(f"\n-> {xlsx}")
    print(f"-> {csv}   (clave ciega: NO abrir antes de anotar)")


if __name__ == "__main__":
    main()
